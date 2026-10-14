# -*- coding: utf-8 -*-
"""sample — 样片选美一键工具(心法 §1.3 的执行力)。

用户说不上来要什么风格时,用同一内容一键生成多风格样片册,
让用户翻页指认——比任何问卷都准。

用法:
    python sample.py --title "你的主题" --subtitle "副题" --styles brutal,dark_premium,editorial
    python sample.py --title "..." --lead "导语" --stats "38项,高危清零;214个,权限回收;95%,脱敏覆盖"

输出:一个 pptx,每个风格 2 页(封面 + 数据页),翻着比即可。
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import build_deck as BD
import pptx_kit as K

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# 每种风格挑两个最有代表性的版式做样片
SAMPLE_PLAN = {
    "brutal": [("cover", "封面:粗黑描边+硬投影+酸性黄"), ("stats", "数据:硬投影大卡")],
    "dark_premium": [("cover", "封面:藏青+黄铜金"), ("stats", "数据:金底数据卡")],
    "editorial": [("ed_cover", "杂志封面:衬线大标题+书脊"), ("ed_figures", "行情数据:头条数字制")],
}


def build_samples(title, subtitle, lead, stats, styles, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = BD.IN(K.PAGE_W), BD.IN(K.PAGE_H)
    blank = prs.slide_layouts[6]
    stats_cards = []
    for i, item in enumerate(stats):
        if "," in item:
            num, label = item.split(",", 1)
            stats_cards.append({"num": num.strip(), "unit": "", "label": label.strip(), "sub": ""})
        else:
            stats_cards.append({"num": item.strip(), "unit": "", "label": "", "sub": ""})
    figures = [{"num": c["num"], "label": c["label"], "sub": ""} for c in stats_cards]

    for style_name in styles:
        if style_name not in SAMPLE_PLAN:
            raise ValueError(f"未知风格 {style_name!r},可选 {sorted(SAMPLE_PLAN)}")
        style = BD.load_style(style_name)
        ctx = {"interactive": False, "style": style}
        for layout, _desc in SAMPLE_PLAN[style_name]:
            if layout == "cover":
                sd = {"type": "cover", "kicker": "SAMPLE · " + style_name.upper(),
                      "title": title, "subtitle": subtitle, "lead": lead,
                      "badge": stats_cards[0]["num"] if stats_cards else "",
                      "ticker": "样片选美 · 指认你偏好的风格"}
            elif layout == "stats":
                sd = {"type": "stats", "title": "核心数据", "cards": stats_cards}
            elif layout == "ed_cover":
                sd = {"type": "ed_cover", "kicker": "SAMPLE · EDITORIAL",
                      "issue": "STYLE SELECTION", "title": title, "size": 54,
                      "spine": "样片选美 · 指认你偏好的风格",
                      "standfirst": lead,
                      "meta": [{"label": "STYLE", "value": style_name},
                               {"label": "LAYOUTS", "value": "杂志三件套"},
                               {"label": "MOOD", "value": "克制 · 文气"}]}
            elif layout == "ed_figures":
                sd = {"type": "ed_figures", "kicker": "THE NUMBERS", "title": "关键数字。",
                      "figures": figures[:3], "body": lead}
            slide = prs.slides.add_slide(blank)
            BD.RENDERERS[layout](slide, sd, ctx)
            # 页角标注风格名,方便用户指认
            K.text(slide, 0.85, 7.06, 5.0, 0.3, f"◀ {style_name} · 样片", size=10,
                   color=BD.col(style, "muted"))

    prs.save(out_path)
    print(f"saved: {out_path} | slides: {len(prs.slides._sldIdLst)} "
          f"(每风格 2 页,顺序: {', '.join(styles)})")


def main():
    ap = argparse.ArgumentParser(description="样片选美:一键生成多风格样片册")
    ap.add_argument("--title", required=True, help="主题/标题")
    ap.add_argument("--subtitle", default="", help="副标题")
    ap.add_argument("--lead", default="一句话导语。", help="导语/正文")
    ap.add_argument("--stats", default="", help="数据,分号分隔,如 '38项,高危清零;214个,权限回收'")
    ap.add_argument("--styles", default="brutal,dark_premium,editorial", help="逗号分隔")
    ap.add_argument("--out", default=r"D:\资料M2\mywork\WorkBuddy\PPT\outputs\style-samples.pptx")
    a = ap.parse_args()
    stats = [s.strip() for s in a.stats.split(";") if s.strip()] or ["01", "02", "03"]
    styles = [s.strip() for s in a.styles.split(",") if s.strip()]
    build_samples(a.title, a.subtitle, a.lead, stats, styles, a.out)


if __name__ == "__main__":
    main()
